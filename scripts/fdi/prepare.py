"""Normalize observations without interpolating gaps or inventing joint distributions."""
from __future__ import annotations
import json, math, re, unicodedata
from pathlib import Path
import numpy as np
import pandas as pd
import openpyxl, pycountry
from .storage import read_json, save_json

def number(v):
    if v is None or isinstance(v,bool):return None
    try:
        x=float(v)
        return x if math.isfinite(x) else None
    except (ValueError,TypeError):return None
def norm(v):return ''.join(c for c in unicodedata.normalize('NFKD',str(v or '').strip().lower()) if not unicodedata.combining(c))

COUNTRIES={'Estados Unidos':'USA','Países Baixos':'NLD','Luxemburgo':'LUX','Espanha':'ESP','França':'FRA','Suíça':'CHE','Reino Unido':'GBR','Alemanha':'DEU','Japão':'JPN','Canadá':'CAN','Itália':'ITA','Bélgica':'BEL','Portugal':'PRT','Uruguai':'URY','Brasil':'BRA','China':'CHN','Hong Kong':'HKG','Hong Kong (China)':'HKG','Ilhas Virgens Britânicas':'VGB','Ilhas Cayman':'CYM','Cayman, Ilhas':'CYM','Bermudas':'BMU','Bahamas':'BHS','Áustria':'AUT','Suécia':'SWE','Noruega':'NOR','Dinamarca':'DNK','Finlândia':'FIN','Irlanda':'IRL','Malta':'MLT','Chipre':'CYP','México':'MEX','Chile':'CHL','Argentina':'ARG','Colômbia':'COL','Peru':'PER','Venezuela':'VEN','Panamá':'PAN','Paraguai':'PRY','Bolívia':'BOL','Equador':'ECU','Singapura':'SGP','Cingapura':'SGP','Índia':'IND','Coreia do Sul':'KOR','Coréia do Sul':'KOR','Coreia (Sul)':'KOR','Austrália':'AUS','Nova Zelândia':'NZL','África do Sul':'ZAF','Emirados Árabes Unidos':'ARE','Israel':'ISR','Rússia':'RUS','Rússia, Federação da':'RUS','Turquia':'TUR','Hungria':'HUN','Polônia':'POL','República Tcheca':'CZE','Tcheca, República':'CZE','Liechtenstein':'LIE','Angola':'AGO','Curaçao':'CUW','Curaçau':'CUW','Taiwan':'TWN','Taiwan (Formosa)':'TWN','Gibraltar':'GIB','Jersey':'JEY','Guernsey':'GGY','Ilha de Man':'IMN','Man, Ilha de':'IMN','Líbano':'LBN','Líbia':'LBY','Costa Rica':'CRI','Guatemala':'GTM','Barbados':'BRB','Ilhas Virgens (Britânicas)':'VGB','Ilhas Virgens (EUA)':'VIR','Antilhas Holandesas':'ANT','São Cristóvão e Névis':'KNA','São Vicente e Granadinas':'VCT','Seicheles':'SYC','Seychelles':'SYC','Maurício':'MUS','Maurício, Ilhas':'MUS','Ilhas Marshall':'MHL','Arábia Saudita':'SAU','Indonésia':'IDN','Malásia':'MYS','Tailândia':'THA','Filipinas':'PHL','Islândia':'ISL'}

def world_bank(raw):
    countries={c['id']:{'id':c['id'],'name':c['name'],'longitude':number(c['longitude']),'latitude':number(c['latitude']),'region':c['region']['value']} for c in raw('wb_countries')[1] if c['region']['id']!='NA'}
    data={}
    for indicator,field in [('BX.KLT.DINV.CD.WD','flow_usd'),('BX.KLT.DINV.WD.GD.ZS','flow_gdp'),('NY.GDP.MKTP.CD','gdp_usd')]:
        response=raw('wb_'+indicator)
        if response[0]['pages']!=1:raise ValueError('World Bank response paginated unexpectedly')
        for item in response[1]:
            country=item['countryiso3code'];year=int(item['date'])
            if country not in countries or not 1996<=year<=2025:continue
            data.setdefault((country,year),{'country':country,'year':year})[field]=number(item['value'])
    return list(countries.values()),list(data.values())

def read_bcb_workbook(path,country_names,states):
    w=openpyxl.load_workbook(path,data_only=True)
    origins=[];sectors=[];positions=[];joint=[];state_measures=[];checks=[]
    def year_columns(s,row):
        return [(j,int(str(c.value)[:4])) for j,c in enumerate(s[row+1]) if re.match(r'^(19|20)\d{2}',str(c.value or '')) and 1995<=int(str(c.value)[:4])<=2025]
    def country(label):
        iso=COUNTRIES.get(label)
        return iso,country_names.get(iso,{}).get('name',label)
    s=w['1'];cols=year_columns(s,3)
    for j,y in cols:
        if y<1996:continue
        positions.append({'year':y,'total_usd':number(s.cell(7,j+1).value)*1e6,'equity_usd':number(s.cell(8,j+1).value)*1e6,'debt_usd':number(s.cell(9,j+1).value)*1e6})
    for sheet,basis in [('5','immediate'),('6','ultimate')]:
        s=w[sheet];cols=year_columns(s,4)
        for row in list(s)[10:]:
            label=str(row[0].value or '').strip()
            if not label or not any(number(row[j].value) is not None for j,y in cols):continue
            iso,english=country(label)
            for j,y in cols:
                value=number(row[j].value)
                total=number(s.cell(10,j+1).value)
                if value is not None:origins.append({'country':iso,'name':english,'original_name':label,'basis':basis,'year':y,'value_usd':value*1e6,'share':value/total if total else None})
        for j,y in cols:
            actual=sum(x['value_usd'] for x in origins if x['basis']==basis and x['year']==y)
            expected=float(s.cell(10,j+1).value)*1e6
            checks.append({'table':sheet,'year':y,'sum':actual,'expected':expected,'difference':actual-expected})
    s=w['10'];cols=year_columns(s,4);parent=''
    sector_ids={}
    for row in list(s)[10:]:
        label=str(row[0].value or '').strip();level=int(row[0].alignment.indent or 0)
        if not label or not any(number(row[j].value) is not None for j,y in cols):continue
        if level==0:parent=label
        key=parent+'/'+label;sector_ids[label]=key
        for j,y in cols:
            value=number(row[j].value);total=number(s.cell(10,j+1).value)
            if value is not None:sectors.append({'id':key,'name':label,'group':parent,'level':level,'year':y,'value_usd':value*1e6,'share':value/total if total else None})
    # Joint tables are repeated year blocks. Retain the actual published cells.
    for sheet,basis in [('13','immediate'),('14','ultimate')]:
        s=w[sheet];year=None;headers={}
        for row in s:
            label=str(row[0].value or '').strip()
            if label.startswith('ANO:'):
                year=int(re.search(r'\d{4}',label).group());headers={j:str(c.value).strip() for j,c in enumerate(row) if j>0 and c.value}
                continue
            if not year or year>2025 or not (re.match(r'^[A-U] - ',label) or norm(label).startswith('outros')):continue
            level=int(row[0].alignment.indent or 0)
            for j,name in headers.items():
                value=number(row[j].value)
                if value is None:continue
                if norm(name)=='total':continue
                iso,english=country(name)
                joint.append({'year':year,'basis':basis,'country':iso,'country_name':english,'sector':label,'level':level,'value_usd':value*1e6})
    state_map={norm(s['nome']):str(s['id']) for s in states};state_map['goias']=state_map.get('goias','52')
    for sheet,y in [('15',2010),('16',2015),('17',2020),('18',2025)]:
        s=w[sheet]
        # Only the industrial allocation is present in these geographic tables.
        if str(s.cell(8,1).value).strip()!='Indústria':raise ValueError('BCB geographic industrial header changed')
        for j,c in enumerate(s[5]):
            sid=state_map.get(norm(c.value))
            if sid is not None and number(s.cell(8,j+1).value) is not None:
                state_measures.append({'state':sid,'year':y,'measure':'industrial_equity_position_usd','value':float(s.cell(8,j+1).value)*1e6,'definition':'BCB inward industrial equity position allocated by fixed-asset location; excludes other industries and external/unallocated locations','table':sheet})
    s=w['19']
    for row in list(s)[7:]:
        sid=state_map.get(norm(row[0].value))
        if sid:
            for j,y in [(1,2015),(2,2020),(3,2025)]:
                v=number(row[j].value)
                if v is not None:state_measures.append({'state':sid,'year':y,'measure':'foreign_firm_revenue_share_pct','value':v*100,'definition':'Share of foreign-invested firms gross revenue, not FDI flows','table':'19'})
    if not joint:raise ValueError('Published country-sector tables could not be parsed; refusing an empty joint export')
    return {'origins':origins,'sectors':sectors,'positions':positions,'joint':joint,'state_measures':state_measures,'reconciliation':checks}

def ibge_records(response):
    for variable in response:
        for result in variable['resultados']:
            category=result.get('classificacoes',[])
            category_id=next(iter(category[0]['categoria'])) if category else ''
            for series in result['series']:
                loc=series['localidade']
                for period,value in series['serie'].items():
                    yield {'id':loc['id'],'name':loc['nome'],'period':period,'variable':str(variable['id']),'category':category_id,'value':number(value)}

def prepare(study):
    manifest=read_json(study.run/'raw/sources.json')
    def path(key):return study.run/'raw'/manifest[key]['file']
    def raw(key):return read_json(path(key))
    inputs=[study.run/'raw/sources.json']+[path(k) for k,v in manifest.items() if v['status']=='observed']
    if study.cached('preparation',inputs):return read_json(study.run/'prepared/descriptive.json')
    countries,annual=world_bank(raw);country_names={c['id']:c for c in countries}
    states=raw('ibge_states');cities=raw('ibge_municipalities')
    bcb=read_bcb_workbook(path('bcb_positions'),country_names,states)
    annual_bcb={}
    monthly={}
    for key in ['sgs_22885','sgs_22891','sgs_22892','sgs_22893','sgs_3698']:
        d=pd.DataFrame(raw(key));d['date']=pd.to_datetime(d['data'],dayfirst=True);d['value']=pd.to_numeric(d['valor']);s=d.set_index('date')['value'].sort_index();monthly[key]=s
    for y in range(1996,2026):
        row={'year':y}
        for key,col in [('sgs_22885','total'),('sgs_22891','equity_ex_reinvested'),('sgs_22892','reinvested'),('sgs_22893','debt')]:
            s=monthly[key];v=s[s.index.year==y];row[col]=float(v.sum())*1e6 if len(v)==12 else None
        if all(row[k] is not None for k in ['total','equity_ex_reinvested','reinvested','debt']):
            row['reconciliation_difference']=row['total']-sum(row[k] for k in ['equity_ex_reinvested','reinvested','debt'])
        annual_bcb[y]=row
    quarterly={}
    for key in ['ibge_quarterly','ibge_quarterly_nominal']:
        for r in ibge_records(raw(key)):
            period=pd.Period(f"{r['period'][:4]}Q{int(r['period'][4:])}",freq='Q')
            col=('gdp_level' if r['category']=='90707' else 'gfcf_level') if key=='ibge_quarterly' else 'gdp_nominal_brl_millions'
            quarterly.setdefault(period,{})[col]=r['value']
    frame=pd.DataFrame.from_dict(quarterly,orient='index').sort_index()
    for key in ['NFCI','GDPC1','FEDFUNDS','RBBRBIS']:
        d=pd.read_csv(path('fred_'+key));d.iloc[:,0]=pd.to_datetime(d.iloc[:,0]);s=pd.Series(pd.to_numeric(d.iloc[:,1],errors='coerce').values,index=pd.DatetimeIndex(d.iloc[:,0]))
        frame[key]=s.groupby(s.index.to_period('Q')).mean()
    fdi=monthly['sgs_22885'];fx=monthly['sgs_3698']
    brl=(fdi*fx).groupby(fdi.index.to_period('Q')).sum(min_count=3)
    frame['fdi_gdp']=100*brl/frame['gdp_nominal_brl_millions']
    frame['nfci']=frame.NFCI
    frame['gdp_growth']=100*np.log(frame.gdp_level).diff()
    frame['gfcf_growth']=100*np.log(frame.gfcf_level).diff()
    frame['reer_change']=100*np.log(frame.RBBRBIS).diff()
    frame['us_gdp_growth']=100*np.log(frame.GDPC1).diff()
    frame['fedfunds_change']=frame.FEDFUNDS.diff()
    frame['quarter']=frame.index.quarter
    model_columns=study.config['variables']+['us_gdp_growth','fedfunds_change','quarter']
    panel=frame[model_columns]
    # No dropping interior gaps or pretending adjacent observations are consecutive quarters.
    observed=panel.dropna()
    if len(observed) and len(pd.period_range(observed.index[0],observed.index[-1],freq='Q'))!=len(observed):raise ValueError('Interior quarterly gap; estimation prohibited')
    observed.index=observed.index.astype(str);observed.index.name='period'
    observed.to_csv(study.run/'prepared/quarterly.csv')
    state_accounts=[];municipal_accounts={}
    def accounts(response):
        grouped={}
        fields={'37':'gdp_brl','498':'gva_brl','513':'agriculture_brl','517':'industry_brl','6575':'services_brl','525':'public_brl'}
        for r in ibge_records(response):
            k=(r['id'],int(r['period']));grouped.setdefault(k,{'id':r['id'],'name':r['name'],'year':int(r['period'])})[fields[r['variable']]]=r['value']*1000 if r['value'] is not None else None
        return list(grouped.values())
    state_accounts=accounts(raw('ibge_state_accounts'))
    for key,record in manifest.items():
        if key.startswith('municipal_accounts_') and len(key.split('_'))==4 and record['status']=='observed':municipal_accounts.setdefault(key.split('_')[2],[]).extend(accounts(raw(key)))
    short_ids={str(c['id'])[:-1]:str(c['id']) for c in cities}
    percap={}
    for key,record in manifest.items():
        if key.startswith('municipal_percap_') and record['status']=='observed':
            for variable in raw(key):
                for r in variable['res']:
                    ident=short_ids.get(str(r['localidade']))
                    if ident:
                        for year,value in r['res'].items():
                            v=number(value)
                            if v is not None:percap[(ident,int(year))]=v
    city_lookup={str(c['id']):{'id':str(c['id']),'name':c['nome'],'state':str(c['id'])[:2]} for c in cities}
    for sid,rows in municipal_accounts.items():
        for r in rows:r['gdp_per_capita']=percap.get((r['id'],r['year']))
        save_json(study.run/'prepared'/f'municipal-{sid}.json',{'municipalities':[v for v in city_lookup.values() if v['state']==sid],'observations':rows})
    geo=raw('world_boundaries');features=[]
    for f in geo['features']:
        p=f['properties'];iso=p.get('ISO_A3_EH') or p.get('ADM0_A3') or p.get('ISO_A3')
        if p.get('CONTINENT')=='Antarctica':continue
        features.append({'type':'Feature','properties':{'name':iso,'label':p.get('NAME_EN') or p.get('NAME'),'iso3':iso},'geometry':f['geometry']})
    save_json(study.run/'prepared/world.geo.json',{'type':'FeatureCollection','features':features})
    geo=raw('brazil_states_geometry')
    for f in geo['features']:f['properties']['name']=str(f['properties']['codarea'])
    save_json(study.run/'prepared/states.geo.json',geo)
    for key,rec in manifest.items():
        if key.startswith('municipal_geometry_') and rec['status']=='observed':
            geo=raw(key)
            for f in geo['features']:f['properties']['name']=str(f['properties']['codarea'])
            save_json(study.run/'prepared'/f"municipal-{key.rsplit('_',1)[1]}.geo.json",geo)
    output={'schema_version':1,'window':[1996,2025],'countries':countries,'annual':annual,'components':list(annual_bcb.values()),'states':[{'id':str(s['id']),'name':s['nome'],'abbreviation':s['sigla']} for s in states],'state_accounts':state_accounts,**bcb,'sources':manifest,'quarterly_coverage':{'start':observed.index[0] if len(observed) else None,'end':observed.index[-1] if len(observed) else None,'observations':len(observed),'missing_periods':len(panel)-len(observed)},'municipal_coverage':{'states':len(municipal_accounts),'municipalities':len(city_lookup),'percap_observations':len(percap),'note':'GDP/GVA reference-2010 series 2002–2023; revised GDP per capita has its own coverage. Sector gaps are not filled.'}}
    save_json(study.run/'prepared/descriptive.json',output)
    study.finish('preparation',list((study.run/'prepared').glob('*')),inputs,details={'annual_observations':len(annual),'quarterly_observations':len(observed),'municipal_states':len(municipal_accounts)})
    return output
